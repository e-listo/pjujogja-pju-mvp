import io
import os
import unittest
from openpyxl import Workbook
from aset_bulk_service import ASSET, LAMP, BatchError, read_file, validate, persist

CODE='PJUP-UH2-26-001'
def sample():
    a=dict(zip(ASSET,[CODE,'PJUP','UH2','Sektor 1','2026','Lokasi uji','-7.8','110.37','Jalan Kota','','Besi','8','Menyala']))
    l=dict(zip(LAMP,[CODE,'LED','60','Uji','2026','Menyala']))
    return [(2,a)],[(2,l),(3,dict(l))]
def checked(a=None,l=None,existing=None):
    x,y=sample()
    return validate(x if a is None else a,y if l is None else l,{'PJUP':1},{'UH2':1},existing or set())

class ValidationTests(unittest.TestCase):
    def test_two_lamps(self):
        a,l=checked();self.assertEqual((len(a),len(l)),(1,2))
    def test_empty_lamps(self):self.assertEqual(len(checked(l=[])[1]),0)
    def test_duplicate(self):
        a,l=sample()
        with self.assertRaises(BatchError):checked(a=a+a)
    def test_existing(self):
        with self.assertRaises(BatchError):checked(existing={CODE})
    def test_bad_coordinate(self):
        for value in ('NaN','100',''):
            a,l=sample();a[0][1]['lat']=value
            with self.assertRaises(BatchError):checked(a=a)
    def test_orphan(self):
        a,l=sample();l[0][1]['kode_aset']='UNKNOWN'
        with self.assertRaises(BatchError):checked(l=l)
    def test_xlsx(self):
        a,l=sample();w=Workbook();w.active.title='Data_Aset';w.active.append(ASSET);w.active.append(list(a[0][1].values()))
        s=w.create_sheet('Data_Lampu');s.append(LAMP)
        for _,r in l:s.append(list(r.values()))
        b=io.BytesIO();w.save(b)
        self.assertEqual(len(read_file(b.getvalue(),'test.xlsx')[1]),2)
    def test_formula(self):
        a,l=sample();w=Workbook();w.active.title='Data_Aset';w.active.append(ASSET);r=list(a[0][1].values());r[6]='=1+1';w.active.append(r);w.create_sheet('Data_Lampu').append(LAMP)
        b=io.BytesIO();w.save(b)
        with self.assertRaises(BatchError):read_file(b.getvalue(),'test.xlsx')

@unittest.skipUnless(os.getenv('PIJAR_DISPOSABLE_DB')=='yes','MariaDB integration opt-in')
class MariaDBTests(unittest.TestCase):
    def setUp(self):
        from flask import Flask
        from sqlalchemy.engine import make_url
        from models import db, KategoriPJU, Wilayah
        uri=os.environ['PIJAR_TEST_DB_URI']
        if make_url(uri).database!='pijar_bulk_test':raise RuntimeError('Only pijar_bulk_test allowed')
        self.app=Flask(__name__);self.app.config.update(SQLALCHEMY_DATABASE_URI=uri,SQLALCHEMY_TRACK_MODIFICATIONS=False)
        db.init_app(self.app);self.ctx=self.app.app_context();self.ctx.push();self.db=db
        db.drop_all();db.create_all()
        db.session.add(KategoriPJU(id=1,kode='PJUP',nama='Uji',aktif=True));db.session.add(Wilayah(id_wilayah=1,kode_wilayah='UH2',nama_kelurahan='Uji',nama_kemantren='Uji'));db.session.commit()
    def tearDown(self):
        self.db.session.rollback();self.db.session.remove();self.db.drop_all();self.ctx.pop()
    def test_save_relations(self):
        from models import AsetPJU, Lampu
        a,l=checked();persist(a,l)
        self.assertEqual(AsetPJU.query.count(),1);self.assertEqual(Lampu.query.count(),2)
        self.assertEqual(AsetPJU.query.one().lampu.count(),2)
    def test_rollback(self):
        from sqlalchemy import event
        from models import AsetPJU, Lampu
        def fail(mapper,connection,target):raise RuntimeError('Injected failure')
        event.listen(Lampu,'before_insert',fail)
        try:
            a,l=checked()
            with self.assertRaises(RuntimeError):persist(a,l)
        finally:event.remove(Lampu,'before_insert',fail)
        self.assertEqual(AsetPJU.query.count(),0);self.assertEqual(Lampu.query.count(),0)

if __name__=='__main__':unittest.main()
