import io
import unittest
from openpyxl import Workbook
from aset_bulk_service import ASSET_V2 as ASSET, LAMP_V2 as LAMP, ASSET_LEGACY, LAMP_LEGACY, BatchError, read_file, validate, code_variants

class FourDigitImportTest(unittest.TestCase):
    def row(self, **changes):
        r = {k: '' for k in ASSET}
        r.update(nomor_urut='0001', kode_kategori='PJUP', kode_wilayah='PA1', sektor='1', tahun_pemasangan='2026', alamat='Alamat uji', lat='-7.8', lng='110.37', kategori_jalan='KOTA', status='Menyala')
        r.update(changes)
        return r

    def check(self, r=None, lamps=None, existing=None):
        return validate([(2, self.row() if r is None else r)], lamps or [], {'PJUP': 1}, {'PA1': 1, 'PA2': 2}, existing or set())

    def test_legacy_positional_headers(self):
        from aset_bulk_service import ASSET as old_asset, LAMP as old_lamp
        self.assertEqual(old_asset, ASSET_LEGACY)
        self.assertEqual(old_lamp, LAMP_LEGACY)
        self.assertNotIn('nomor_urut', old_asset)
        self.assertEqual(old_asset[1:3], ['kode_kategori', 'kode_wilayah'])
        r = self.row(kode_aset='PJUP-PA1-26-001', nomor_urut='')
        positional = dict(zip(old_asset, [r[k] for k in old_asset]))
        self.assertEqual(self.check(positional)[0][0]['kode_aset'], 'PJUP-PA1-26-001')

    def test_padding(self):
        for n in ('1', '001', '0001', '9999'):
            with self.subTest(n=n):
                assets, _ = self.check(self.row(nomor_urut=n))
                self.assertEqual(assets[0]['kode_aset'], f'PJUP-PA1-26-{int(n):04d}')

    def test_invalid_number(self):
        for n in ('', '0', '0000', '-1', '10000', '1.5', '1e2', 'abcd'):
            with self.subTest(n=n), self.assertRaises(BatchError):
                self.check(self.row(nomor_urut=n))

    def test_aliases(self):
        assets, _ = self.check()
        self.assertEqual(assets[0]['sektor'], 'Sektor 1')
        self.assertEqual(assets[0]['kategori_jalan'], 'Jalan Kota')

    def test_old_code_preserved(self):
        assets, _ = self.check(self.row(kode_aset='PJUP-PA1-26-001', nomor_urut=''))
        self.assertEqual(assets[0]['kode_aset'], 'PJUP-PA1-26-001')

    def test_component_conflicts(self):
        for change in ({'nomor_urut': '2'}, {'tahun_pemasangan': '2025'}, {'kode_wilayah': 'PA2'}):
            with self.subTest(change=change), self.assertRaises(BatchError):
                self.check(self.row(kode_aset='PJUP-PA1-26-0001', **change))

    def test_existing_equivalent(self):
        for code in ('PJUP-PA1-26-001', 'PJUP-PA1-26-0001'):
            with self.subTest(code=code), self.assertRaises(BatchError):
                self.check(existing={code})

    def test_batch_equivalent(self):
        with self.assertRaises(BatchError):
            validate([(2, self.row()), (3, self.row(kode_aset='PJUP-PA1-26-001', nomor_urut=''))], [], {'PJUP': 1}, {'PA1': 1}, set())

    def test_same_number_other_region(self):
        assets, _ = validate([(2, self.row()), (3, self.row(kode_wilayah='PA2'))], [], {'PJUP': 1}, {'PA1': 1, 'PA2': 2}, set())
        self.assertEqual(len(assets), 2)

    def test_missing_address(self):
        with self.assertRaisesRegex(BatchError, 'Alamat wajib'):
            self.check(self.row(alamat=''))

    def test_lamp_components_and_old_alias(self):
        for fields in ({'kode_aset': 'PJUP-PA1-26-001'}, {'nomor_urut': '1', 'kode_kategori': 'PJUP', 'kode_wilayah': 'PA1', 'tahun_pemasangan': '2026'}):
            r = {k: '' for k in LAMP}; r.update(jenis_lampu='LED', daya_watt='50', status_lampu='Menyala', **fields)
            _, lights = self.check(lamps=[(2, r)])
            self.assertEqual(lights[0]['kode_aset'], 'PJUP-PA1-26-0001')

    def test_unlinked_lamp(self):
        r = {k: '' for k in LAMP}; r.update(kode_aset='PJUP-PA2-26-0001', jenis_lampu='LED', status_lampu='Menyala')
        with self.assertRaisesRegex(BatchError, 'tidak terhubung'):
            self.check(lamps=[(2, r)])

    def test_both_workbook_versions(self):
        for headers, lamp_headers in ((ASSET, LAMP), (ASSET_LEGACY, LAMP_LEGACY)):
            w = Workbook(); w.active.title = 'Data_Aset'; w.active.append(headers)
            r = self.row()
            if headers == ASSET_LEGACY:
                r['kode_aset'] = 'PJUP-PA1-26-001'
            w.active.append([r[k] for k in headers]); w.create_sheet('Data_Lampu').append(lamp_headers)
            b = io.BytesIO(); w.save(b); w.close()
            assets, lamps = read_file(b.getvalue(), 'test.xlsx')
            self.assertEqual(len(self.check(assets[0][1], lamps)[0]), 1)

    def test_formula_rejected(self):
        w = Workbook(); w.active.title = 'Data_Aset'; w.active.append(ASSET)
        w.active.append(['=1+1']); w.create_sheet('Data_Lampu').append(LAMP)
        b = io.BytesIO(); w.save(b); w.close()
        with self.assertRaisesRegex(BatchError, 'Formula'):
            read_file(b.getvalue(), 'test.xlsx')

    def test_variants(self):
        self.assertEqual(code_variants('PJUP-PA1-26-0001'), {'PJUP-PA1-26-001', 'PJUP-PA1-26-0001'})

if __name__ == '__main__':
    unittest.main()
