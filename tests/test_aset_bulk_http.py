import io
import os
import unittest
from types import SimpleNamespace
from openpyxl import Workbook, load_workbook

@unittest.skipUnless(os.getenv('PIJAR_DISPOSABLE_DB')=='yes','Database uji belum tersedia')
class HTTPTests(unittest.TestCase):
    def setUp(self):
        from flask import Flask
        from sqlalchemy.engine import make_url
        from werkzeug.security import generate_password_hash
        from models import db, Pengguna, KategoriPJU, Wilayah
        from auth_routes import auth_bp, _BLACKLIST
        from aset_bulk_routes import bp
        uri=os.environ['PIJAR_TEST_DB_URI']
        if make_url(uri).database!='pijar_bulk_test':raise RuntimeError('Database uji wajib pijar_bulk_test')
        self.app=Flask(__name__);self.app.config.update(TESTING=True,JWT_SECRET='ci-only-not-production-secret-123456789',JWT_EXP_HOURS=1,MAX_CONTENT_LENGTH=5*1024*1024,SQLALCHEMY_DATABASE_URI=uri,SQLALCHEMY_TRACK_MODIFICATIONS=False)
        db.init_app(self.app);self.app.register_blueprint(auth_bp);self.app.register_blueprint(bp)
        self.ctx=self.app.app_context();self.ctx.push();self.db=db;_BLACKLIST.clear()
        self.addCleanup(self.cleanup)
        db.drop_all();db.create_all()
        for role in ('admin','koordinator','teknisi'):
            db.session.add(Pengguna(nama_lengkap=role,username=role,password_hash=generate_password_hash('ci-only'),peran=role,status_aktif=True))
        db.session.add(KategoriPJU(id=1,kode='PJUP',nama='Uji',aktif=True));db.session.add(Wilayah(id_wilayah=1,kode_wilayah='UH2',nama_kelurahan='Uji',nama_kemantren='Uji'));db.session.commit()
        self.client=self.app.test_client()
        self._xlsx=None
    def cleanup(self):
        self.db.session.rollback();self.db.session.remove();self.db.drop_all();self.ctx.pop()
    def headers(self,role='admin'):
        if role=='regu':
            from auth_routes import _buat_token
            token,_=_buat_token(SimpleNamespace(id_pengguna=99,username='test-regu',peran='regu',nama_lengkap='Uji regu'))
            return {'Authorization':'Bearer '+token}
        r=self.client.post('/api/auth/login',json={'username':role,'password':'ci-only'})
        self.assertEqual(r.status_code,200)
        return {'Authorization':'Bearer '+r.get_json()['token']}
    def data(self):
        # Dibuat sekali per tes: openpyxl menyimpan waktu pembuatan (detik) di dalam berkas,
        # sehingga berkas yang dibuat ulang bisa berbeda byte dan hash pratinjau tidak cocok saat commit.
        if self._xlsx is None:
            self._xlsx=self.bangun_xlsx()
        return self._xlsx
    def bangun_xlsx(self):
        from aset_bulk_service import ASSET,LAMP
        w=Workbook();w.active.title='Data_Aset';w.active.append(ASSET)
        w.active.append(['PJUP-UH2-26-001','PJUP','UH2','Sektor 1',2026,'Uji',-7.8,110.37,'Jalan Kota','','Besi',8,'Menyala'])
        s=w.create_sheet('Data_Lampu');s.append(LAMP)
        for watt in (60,90):s.append(['PJUP-UH2-26-001','LED',watt,'Uji',2026,'Menyala'])
        b=io.BytesIO();w.save(b);return b.getvalue()
    def post(self,mode,headers,**fields):
        return self.client.post('/api/aset/import/'+mode,data={'file':(io.BytesIO(self.data()),'uji.xlsx'),**fields},headers=headers)
    def test_missing_token(self):
        self.assertEqual(self.post('preview',{}).status_code,401)
    def test_invalid_token(self):
        self.assertEqual(self.post('preview',{'Authorization':'Bearer invalid'}).status_code,401)
    def test_forbidden_roles(self):
        for role in ('teknisi',):
            with self.subTest(role=role):
                h=self.headers(role)
                self.assertEqual(self.post('preview',h).status_code,403)
                self.assertEqual(self.post('commit',h).status_code,403)
                self.assertEqual(self.client.get('/api/aset/import/template',headers=h).status_code,403)
    def test_token_akun_tidak_terdaftar_ditolak(self):
        h=self.headers('regu')
        self.assertEqual(self.post('preview',h).status_code,401)
        self.assertEqual(self.post('commit',h).status_code,401)
        self.assertEqual(self.client.get('/api/aset/import/template',headers=h).status_code,401)
    def test_akun_nonaktif_ditolak(self):
        from models import Pengguna
        h=self.headers('koordinator')
        self.assertEqual(self.post('preview',h).status_code,200)
        Pengguna.query.filter_by(username='koordinator').update({'status_aktif':False});self.db.session.commit()
        self.assertEqual(self.post('preview',h).status_code,401)
    def test_koordinator_preview(self):
        from models import AsetPJU
        r=self.post('preview',self.headers('koordinator'));self.assertEqual(r.status_code,200)
        self.assertEqual(r.get_json()['total_lampu'],2);self.assertEqual(AsetPJU.query.count(),0)
    def test_commit(self):
        from models import AsetPJU
        h=self.headers();r=self.post('preview',h);self.assertEqual(r.status_code,200)
        saved=self.post('commit',h,confirm='yes',sha256=r.get_json()['sha256'])
        self.assertEqual(saved.status_code,201);self.assertEqual(AsetPJU.query.one().lampu.count(),2)
    def test_confirmation_required(self):
        self.assertEqual(self.post('commit',self.headers()).status_code,400)
    def test_hash_mismatch(self):
        from models import AsetPJU
        self.assertEqual(self.post('commit',self.headers(),confirm='yes',sha256='wrong').status_code,400)
        self.assertEqual(AsetPJU.query.count(),0)
    def test_template(self):
        r=self.client.get('/api/aset/import/template',headers=self.headers());self.assertEqual(r.status_code,200)
        w=load_workbook(io.BytesIO(r.data),read_only=True);self.assertIn('Data_Lampu',w.sheetnames);w.close()
    def test_logged_out_token(self):
        h=self.headers();self.assertEqual(self.client.post('/api/auth/logout',headers=h).status_code,200)
        self.assertEqual(self.post('preview',h).status_code,401)
    def test_corrupt_file(self):
        r=self.client.post('/api/aset/import/preview',data={'file':(io.BytesIO(b'broken'),'bad.xlsx')},headers=self.headers())
        self.assertEqual(r.status_code,400)

if __name__=='__main__':unittest.main()
