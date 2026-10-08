import csv
import io
import hashlib
import zipfile
from flask import Blueprint, request, jsonify, Response, current_app
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError
from models import db, AsetPJU, KategoriPJU, Wilayah
from auth_routes import role_required
from aset_bulk_service import ASSET_V2 as ASSET, LAMP_V2 as LAMP, BatchError, read_file, validate, persist, normalize_assets, code_variants

bp = Blueprint('aset_bulk', __name__)
MAX_FILE = 4 * 1024 * 1024

def attachment(data, filename, mimetype):
    return Response(data, mimetype=mimetype, headers={'Content-Disposition': f'attachment; filename="{filename}"', 'Content-Length': str(len(data)), 'Cache-Control': 'no-store'})

def prepare():
    f = request.files.get('file')
    if not f or not f.filename:
        raise BatchError('Pilih berkas')
    raw = f.stream.read(MAX_FILE + 1)
    if len(raw) > MAX_FILE:
        raise BatchError('Maksimum 4 MB')
    try:
        assets, lamps = read_file(raw, f.filename)
    except BatchError:
        raise
    except (ValueError, UnicodeError, zipfile.BadZipFile, KeyError) as e:
        raise BatchError('Berkas tidak sesuai template') from e
    assets = normalize_assets(assets)
    categories = {k.kode: k.id for k in KategoriPJU.query.filter_by(aktif=True).all()}
    regions = {w.kode_wilayah: w.id_wilayah for w in Wilayah.query.all()}
    codes = sorted({c for _, r in assets for c in code_variants(r['kode_aset'])})
    existing = set()
    for start in range(0, len(codes), 500):
        existing.update(x[0] for x in db.session.query(AsetPJU.kode_aset).filter(AsetPJU.kode_aset.in_(codes[start:start + 500])).all())
    assets, lamps = validate(assets, lamps, categories, regions, existing)
    return assets, lamps, hashlib.sha256(raw).hexdigest()

@bp.errorhandler(BatchError)
def invalid(e):
    db.session.rollback()
    return jsonify(success=False, error=str(e)), 400

@bp.route('/api/aset/import/preview', methods=['POST'])
@role_required('admin', 'koordinator')
def preview():
    assets, lamps, digest = prepare()
    return jsonify(success=True, valid=True, sha256=digest, total_aset=len(assets), total_lampu=len(lamps), preview_aset=assets[:20], preview_lampu=lamps[:20])

@bp.route('/api/aset/import/commit', methods=['POST'])
@role_required('admin', 'koordinator')
def commit():
    if request.form.get('confirm') != 'yes':
        raise BatchError('Konfirmasi diperlukan')
    assets, lamps, digest = prepare()
    if request.form.get('sha256') != digest:
        raise BatchError('Berkas berbeda dari pratinjau')
    try:
        return jsonify(success=True, **persist(assets, lamps)), 201
    except IntegrityError:
        return jsonify(success=False, error='Konflik data; seluruh batch dibatalkan'), 409
    except Exception:
        current_app.logger.exception('Bulk import gagal')
        return jsonify(success=False, error='Penyimpanan gagal; seluruh batch dibatalkan'), 500

@bp.route('/api/aset/import/template', methods=['GET'])
@role_required('admin', 'koordinator')
def template():
    fmt = request.args.get('format', 'xlsx')
    if fmt == 'csv':
        s = io.StringIO(); csv.writer(s).writerow(ASSET)
        return attachment(s.getvalue().encode('utf-8-sig'), 'template_aset.csv', 'text/csv')
    if fmt != 'xlsx':
        raise BatchError('Format tidak didukung')
    w = Workbook(); w.active.title = 'Data_Aset'; w.active.append(ASSET)
    w.create_sheet('Data_Lampu').append(LAMP)
    refs = w.create_sheet('Referensi'); refs.append(['kode_kategori', 'nama'])
    for k in KategoriPJU.query.filter_by(aktif=True).order_by(KategoriPJU.kode).all():
        refs.append([k.kode, k.nama])
    refs.append([]); refs.append(['kode_wilayah', 'kelurahan', 'kemantren'])
    for r in Wilayah.query.order_by(Wilayah.kode_wilayah).all():
        refs.append([r.kode_wilayah, r.nama_kelurahan, r.nama_kemantren])
    for name, headers in [('Data_Aset', ASSET), ('Data_Lampu', LAMP)]:
        w[name].freeze_panes = 'A2'
        for row in w[name].iter_rows(min_row=2, max_row=101, max_col=len(headers)):
            for c in row:
                if headers[c.column - 1] in ('kode_aset', 'nomor_urut', 'kode_kategori', 'kode_wilayah'):
                    c.number_format = '@'
    g = w.create_sheet('Panduan')
    for text in [
        'Satu baris satu aset; satu baris satu lampu. Template lama tetap didukung.',
        'Isi kode_aset lengkap ATAU nomor_urut + kode_kategori + kode_wilayah + tahun_pemasangan.',
        'Nomor manual 1-9999; 1/001/0001 menjadi 0001. Contoh PJUP-PA1-26-0001. Nomor 0000 dilarang.',
        'Jika kode_aset dan komponen diisi bersama, semuanya harus sesuai; kode lama 3 digit tidak diubah.',
        'Lampu: gunakan kode_aset lengkap atau komponen yang sama dengan aset batch. Jangan nomor saja.',
        'Padanan kode 3/4 digit dianggap duplikat; data yang ada tidak ditimpa.',
        'Sektor: Sektor 1-4 atau angka 1-4. Kategori jalan: Jalan Kota/Jalan Lingkungan/Jalan Lingkungan Kampung/Lainnya; KOTA diterima.',
        'Alamat wajib. Wilayah harus diperiksa pengguna; sistem tidak menebak wilayah dari alamat.',
        'Maksimum 4 MB, 1000 baris total. Formula dilarang. Sheet lampu boleh kosong.',
        'Gunakan Excel. CSV belum mendukung pemisah titik koma.',
    ]:
        g.append([text])
    b = io.BytesIO(); w.save(b); w.close()
    return attachment(b.getvalue(), 'template_import_pijar.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
