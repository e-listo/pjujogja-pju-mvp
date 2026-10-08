"""Bulk import validation and atomic persistence."""
import csv
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

ASSET = 'kode_aset kode_kategori kode_wilayah sektor tahun_pemasangan alamat lat lng kategori_jalan sub_kategori_lainnya jenis_tiang tinggi_meter status'.split()
LAMP = 'kode_aset jenis_lampu daya_watt merk tahun_pasang status_lampu'.split()

class BatchError(ValueError):
    pass

def rows(values, headers):
    it = iter(values)
    if list(next(it, [])) != headers:
        raise BatchError('Header tidak sesuai template')
    result = []
    for line, raw in enumerate(it, 2):
        raw = list(raw)
        if len(raw) > len(headers):
            raise BatchError('Kolom tambahan tidak diizinkan')
        raw += [None] * (len(headers) - len(raw))
        r = {k: '' if v is None else str(v).strip() for k, v in zip(headers, raw)}
        if any(r.values()):
            result.append((line, r))
        if len(result) > 1000:
            raise BatchError('Maksimum 1000 baris')
    return result

def read_file(data, filename):
    if len(data) > 5 * 1024 * 1024:
        raise BatchError('Maksimum file 5 MB')
    if filename.lower().endswith('.csv'):
        return rows(csv.reader(io.StringIO(data.decode('utf-8-sig'))), ASSET), []
    if not filename.lower().endswith('.xlsx'):
        raise BatchError('Gunakan XLSX atau CSV')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if len(z.infolist()) > 200 or sum(x.file_size for x in z.infolist()) > 30 * 1024 * 1024:
            raise BatchError('Isi Excel terlalu besar')
    w = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if not {'Data_Aset', 'Data_Lampu'}.issubset(w.sheetnames):
            raise BatchError('Dua sheet wajib tersedia')
        for name, headers in [('Data_Aset', ASSET), ('Data_Lampu', LAMP)]:
            s = w[name]
            if s.max_row > 1001 or s.max_column > len(headers):
                raise BatchError('Dimensi sheet melebihi batas')
            for row in s.iter_rows():
                if any(c.data_type == 'f' for c in row):
                    raise BatchError('Formula tidak diizinkan')
        return rows(w['Data_Aset'].values, ASSET), rows(w['Data_Lampu'].values, LAMP)
    finally:
        w.close()

def number(v, low, high, integer=False, places=None, label='Angka'):
    if v == '':
        return None
    try:
        d = Decimal(v.replace(',', '.'))
    except InvalidOperation:
        raise BatchError(f'{label} bukan angka')
    if not d.is_finite():
        raise BatchError(f'{label} bukan angka')
    if not low <= d <= high:
        raise BatchError(f'{label} harus antara {low} dan {high}')
    if integer and d != d.to_integral_value():
        raise BatchError(f'{label} harus bilangan bulat')
    if places is not None and d != d.quantize(Decimal(1).scaleb(-places)):
        raise BatchError(f'{label} maksimum {places} angka desimal')
    return int(d) if integer else d

def validate(assets, lamps, categories, regions, existing):
    if not assets or len(assets) + len(lamps) > 1000:
        raise BatchError('Isi 1 sampai 1000 baris total')
    output = []; lights = []; seen = set()
    for line, r in assets:
        try:
            c = r['kode_aset'].upper(); k = r['kode_kategori'].upper();w = r['kode_wilayah'].upper()
            if not re.fullmatch(r'[A-Z]{3,6}-[A-Z]{2}\d-\d{2}-\d{3}', c):
                raise BatchError('Format kode salah')
            if c in seen or c in existing:
                raise BatchError('Kode duplikat; tidak menimpa')
            seen.add(c)
            if k not in categories or w not in regions:
                raise BatchError('Master kategori aktif/wilayah tidak ditemukan')
            if c.split('-')[:2] != [k, w]:
                raise BatchError('Kode tidak sesuai kategori/wilayah')
            if not r['alamat'] or len(r['alamat']) > 255:
                raise BatchError('Alamat wajib, maksimum 255 karakter')
            if r['sektor'] not in ('', 'Sektor 1', 'Sektor 2', 'Sektor 3', 'Sektor 4'):
                raise BatchError('Sektor salah')
            if r['status'] not in ('Menyala', 'Rusak', 'Dalam Pengerjaan'):
                raise BatchError('Status salah')
            if r['kategori_jalan'] not in ('Jalan Kota', 'Jalan Lingkungan', 'Jalan Lingkungan Kampung', 'Lainnya'):
                raise BatchError('Kategori jalan salah')
            sub = r['sub_kategori_lainnya']
            if r['kategori_jalan'] == 'Lainnya':
                if sub not in ('Taman', 'Makam', 'Sorot Sungai', 'Hias/Budaya'):
                    raise BatchError('Subkategori wajib')
            elif sub:
                raise BatchError('Subkategori hanya untuk Lainnya')
            if len(r['jenis_tiang']) > 50:
                raise BatchError('Jenis tiang terlalu panjang')
            lat = number(r['lat'], -90, 90, places=8, label='lat'); lng = number(r['lng'], -180, 180, places=8, label='lng')
            if lat is None or lng is None:
                raise BatchError('Koordinat wajib')
            year = number(r['tahun_pemasangan'], 1901, 2155, True, label='tahun_pemasangan')
            if year is not None and str(year)[-2:].zfill(2) != c.split('-')[2]:
                raise BatchError('Tahun tidak sesuai kode')
            output.append(dict(kode_aset=c, id_kategori=categories[k], id_wilayah=regions[w], alamat=r['alamat'], sektor=r['sektor'] or None, tahun_pemasangan=year, lokasi_lat=lat, lokasi_lng=lng, kategori_jalan=r['kategori_jalan'], sub_kategori_lainnya=sub or None, jenis_tiang=r['jenis_tiang'] or None, tinggi_meter=number(r['tinggi_meter'], Decimal('0.1'),Decimal('999.9'), places=1, label='tinggi_meter'), status=r['status']))
        except BatchError as e:
            raise BatchError(f'Data_Aset baris {line}: {e}') from e
    for line, r in lamps:
        try:
            c = r['kode_aset'].upper()
            if c not in seen:
                raise BatchError('Lampu tidak terhubung ke aset batch ini')
            if r['status_lampu'] not in ('Menyala', 'Mati', 'Redup', 'Rusak'):
                raise BatchError('Status lampu salah')
            kind = r['jenis_lampu'].strip()
            if not kind:
                raise BatchError('Jenis lampu wajib')
            if len(kind) > 50 or len(r['merk']) > 100:
                raise BatchError('Jenis maksimum 50 dan merek maksimum 100 karakter')
            lights.append(dict(kode_aset=c, jenis_lampu=kind, merk=r['merk'] or None, daya_watt=number(r['daya_watt'], 1, 32767, True, label='daya_watt'), tahun_pasang=number(r['tahun_pasang'], 1901, 2155, True, label='tahun_pasang'), status_lampu=r['status_lampu']))
        except BatchError as e:
            raise BatchError(f'Data_Lampu baris {line}: {e}') from e
    return output, lights

def persist(assets, lamps):
    from models import db, AsetPJU, Lampu
    try:
        objects = {}
        for values in assets:
            obj = AsetPJU(**values); db.session.add(obj); objects[obj.kode_aset] = obj
        db.session.flush()
        for r in lamps:
            values = dict(r); code = values.pop('kode_aset')
            db.session.add(Lampu(id_aset=objects[code].id_aset, **values))
        db.session.commit()
        return {'total_aset': len(assets), 'total_lampu': len(lamps)}
    except Exception:
        db.session.rollback()
        raise
